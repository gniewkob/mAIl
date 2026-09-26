"""Build data/golden/golden_labels.xlsx (dropdown per row) from golden_labels.csv.
mAIl's own category is NOT shown, so labels are not anchored to the current model.

  cd ~/Repos/priv/mAIl && uv run --with openpyxl python ops/golden_xlsx.py
"""
import csv
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

CATS = ["appointment", "question", "complaint", "other", "billing", "system", "newsletter", "offer", "spam"]
HELP = {
    "appointment": "wizyta / rezerwacja / zmiana terminu (od człowieka lub Booksy)",
    "question": "pytanie klienta lub osoby prywatnej, oczekuje odpowiedzi",
    "complaint": "reklamacja / skarga od człowieka",
    "other": "poczta od człowieka, która nie pasuje wyżej",
    "billing": "faktura, rachunek, płatność, wyciąg, bank",
    "system": "powiadomienie automatyczne: konto, zamówienie, dostawa, bezpieczeństwo",
    "newsletter": "masowa wysyłka, promocja sklepu do klienta",
    "offer": "oferta handlowa B2B / cold outreach / SEO / współpraca",
    "spam": "spam, oszustwo, phishing",
}
src = Path("data/golden/golden_labels.csv")
rows = list(csv.DictReader(src.open(encoding="utf-8")))
wb = Workbook()
ws = wb.active
ws.title = "do oznaczenia"
cols = ["id", "your_label", "from", "subject", "snippet", "notes"]
ws.append(cols)
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

clean = lambda v: ILLEGAL_CHARACTERS_RE.sub("", v or "")  # control chars in spam subjects
for r in rows:
    ws.append([r["id"], "", clean(r["from"]), clean(r["subject"]), clean(r["snippet"]), ""])
dv = DataValidation(type="list", formula1='"' + ",".join(CATS) + '"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"B2:B{len(rows) + 1}")
for c, w in zip("ABCDEF", (14, 14, 34, 50, 90, 30)):
    ws.column_dimensions[c].width = w
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="DDDDDD")
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    row[1].fill = PatternFill("solid", fgColor="FFF2CC")
ws.freeze_panes = "C2"
legend = wb.create_sheet("kategorie")
legend.append(["kategoria", "kiedy"])
for k, v in HELP.items():
    legend.append([k, v])
legend.column_dimensions["A"].width = 14
legend.column_dimensions["B"].width = 80
out = Path("data/golden/golden_labels.xlsx")
wb.save(out)
os.chmod(out, 0o600)
print(out, len(rows))
